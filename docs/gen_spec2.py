# -*- coding: utf-8 -*-
"""
思辩星球 DebateSphere · 产品方案与创新说明书 生成器
内容模型 CONTENT → Word (.docx) + 打印友好 HTML
五个主章节：核心思路 / 功能架构 / 创新能力 / 技术方案 / 知乎生态契合度
"""
import os, io, html as H
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT_DIR = os.path.dirname(os.path.abspath(__file__))
DOCX_PATH = os.path.join(OUT_DIR, "产品方案与创新说明书-思辩星球×知乎生态.docx")
HTML_PATH = os.path.join(OUT_DIR, "产品方案与创新说明书-思辩星球×知乎生态.html")

CN = '微软雅黑'
EN = 'Segoe UI'

INK   = RGBColor(0x1E, 0x29, 0x3B)
INK2  = RGBColor(0x47, 0x55, 0x69)
MUTED = RGBColor(0x64, 0x74, 0x8B)
PRIMARY = RGBColor(0x0F, 0x76, 0x6E)   # teal-700
ACCENT2 = RGBColor(0xB4, 0x53, 0x09)   # amber-700
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

H_FILL = '134E4A'   # table header
Z_FILL = 'F0FDFA'
BORDER = 'CBD5E1'
H1_LINE = '0F766E'
CALLOUT_FILL = 'FFFBEB'
CALLOUT_BORDER = 'FDE68A'
CALLOUT_TEXT = RGBColor(0x92, 0x40, 0x0E)


# ============================================================
# CONTENT MODEL
# ============================================================
CONTENT = [
    ('cover',),

    # ---------------- 摘要 ----------------
    ('h1', '摘要'),
    ('p', '思辩星球（DebateSphere）是一款以 MBTI 人格为身份入口、以 AI 多智能体为辩手与裁判、以实时辩论为核心交互的在线辩论平台。本说明书系统阐述产品的核心思路、功能架构、创新能力、技术方案，以及其与知乎生态的契合度。'),
    ('table', ['维度', '一句话结论'], [
        ['核心思路', '人格降低表达门槛、辩论生产优质内容、AI 补齐对手与裁判——把"敢不敢说"变成"随时能辩"'],
        ['功能架构', '八大系统，覆盖测评、对战、场景、AI、知识、社区、成长与多端'],
        ['创新能力', '三项模式创新 + 五项体验创新 + 三项技术创新，构成可验证的差异化壁垒'],
        ['技术方案', '全栈自研；服务器权威实时对战；知乎直答作为统一 AI 大脑'],
        ['生态契合', '与知乎同源于"理性讨论"，双向供给话题与内容，综合契合度自评 4.3 / 5'],
    ]),
    ('pagebreak',),

    # ---------------- 一、核心思路 ----------------
    ('h1', '一、核心思路'),
    ('h2', '1.1 产品愿景'),
    ('p', '让每一种人格，都成为一颗星球——一个让"表达"不再有门槛、让"分歧"能产出价值、让"思考"被看见的地方。'),
    ('h2', '1.2 出发点：用户的三个真实困境'),
    ('table', ['困境', '具体表现', '造成后果'], [
        ['不敢说', '公开表达有社交压力，怕被喷、怕不专业、怕被贴标签', '大量观点从未被说出口'],
        ['没人吵', '想认真辩一辩，身边却找不到旗鼓相当的对手', '想法无法被检验、无法被推进'],
        ['留不下', '群聊、评论区吵完就散，没有沉淀', '交锋归零，用户没有成长'],
    ]),
    ('h2', '1.3 三条设计回应（核心支点）'),
    ('table', ['支点', '设计回应', '解决的问题'], [
        ['人格即身份', '用 MBTI 人格作为化身，承担"人格立场"而非"个人立场"', '不敢说'],
        ['辩论即内容', '把立论/反驳/总结/评分/金句结构化落库，产出可传播内容', '留不下'],
        ['AI 即队友', 'AI 当辩手、当对手、当裁判、当教练，永远在线', '没人吵'],
    ]),
    ('p', '其中"AI 即队友"是本产品的供给核心：真人对手的供给永远不稳定，而 AI 对手可以做到 7×24 在线、难度自适应、且越辩越强。这从根本上解决了辩论类产品最难的冷启动问题。'),
    ('h2', '1.4 价值主张'),
    ('table', ['面向对象', '核心价值'], [
        ['普通用户', '零压力地表达、被认真回应、看见自己的成长'],
        ['表达欲强的人', '随时可开局的舞台与会认真反驳的对手'],
        ['学习者', '练逻辑、练表达、练口语的实时陪练与训练计划'],
        ['内容生态（知乎）', '源源不断的真实话题与高质量观点内容'],
    ]),
    ('pagebreak',),

    # ---------------- 二、功能架构 ----------------
    ('h1', '二、功能架构'),
    ('p', '产品由八大系统构成，各系统边界清晰、可独立迭代。下表为功能地图总览，其后逐系统展开。'),
    ('h2', '2.1 功能地图总览'),
    ('table', ['系统', '核心功能', '一句话价值'], [
        ['人格系统', '16 型测评 / 五维驱力 / 四层引擎 / 人格演化', '定义"我是谁"'],
        ['对战系统', 'AI 辩论 / 1v1 对话 / PK 实时对战 / 宠物战斗', '定义"怎么辩"'],
        ['场景系统', '法律辩论 / 圆桌会谈 / 脱口秀 / 街头对谈', '定义"在哪辩、辩什么"'],
        ['AI 能力系统', '直答三档 / 四类角色 / 风格学习 / 节奏引擎', '定义"谁来陪"'],
        ['知识系统', '本地 RAG / 云端知识库 / 每日新闻 / 知乎源', '定义"拿什么辩"'],
        ['社区系统', '观点广场 / 评论点赞 / 内容自删', '定义"辩给谁看"'],
        ['成长系统', '七维评分 / 辩论报告 / 训练计划 / 实力分', '定义"辩完得到什么"'],
        ['多端系统', 'Web / Android / Windows / 语音 / 音效', '定义"在哪都能辩"'],
    ]),
    ('h2', '2.2 人格系统'),
    ('table', ['功能', '说明'], [
        ['16 型 MBTI 完整测评', '标准四维判定，输出人格类型与画像'],
        ['五维驱力刻画', '认知 / 情感 / 意志 / 社交 / 表达五个驱力维度'],
        ['四层人格引擎', '画像层 → 认知层 → 演化层 → 记忆层'],
        ['人格演化', '人格随辩论历史确定性演化，是"活的"而非一次性标签'],
        ['人格色身份体系', '16 色光谱 + 头像 + 光晕，构成完整视觉身份'],
    ]),
    ('h2', '2.3 对战系统'),
    ('table', ['功能', '说明'], [
        ['AI 多智能体辩论', '审题 → 检索 → 立场 → 发言 → 裁判的完整流程，思考链折叠 + 流式输出'],
        ['1v1 对话模式', '与单个 AI 人格自由对话，支持自由选题'],
        ['PK 实时对战', '房间制实时辩论：立论 → 自由辩论 → 总结 → AI 裁判评分'],
        ['AI 三档对手', '初级 / 中级 / 大师，难度自适应'],
        ['逐句发言节奏引擎', 'AI 发言逐句连发、观点表达完整后再切阶段'],
        ['回合控制', '退出房间、请 AI 继续说（poke）、AI 主动开场（kickoff）'],
        ['服务器权威宠物战斗', '发言即攻击，伤害由服务端计算并广播，两端一致'],
        ['宠物养成体系', '8 种像素宠物 + 装备加成 + 积分商城 + 升级系统'],
    ]),
    ('h2', '2.4 场景系统'),
    ('table', ['场景', '模式', '玩法特征', 'AI 命题来源'], [
        ['法律辩论 · 普法现场', '对抗', '法官 + 控辩双方，落法律条文与判例', 'AI 结合热榜生成法律议题'],
        ['圆桌会谈 · 哲思夜话', '自由（不分胜负）', '无正反方，把一个问题聊透', 'AI 生成人生哲学命题'],
        ['脱口秀大会', '自由', '幽默是唯一正义，必须落到生活细节', 'AI 生成生活向话题'],
        ['街头对谈 · 人间观察', '自由', '说人话、不端着，随意真实', 'AI 生成日常困惑类话题'],
    ]),
    ('h2', '2.5 AI 能力系统'),
    ('table', ['功能', '说明'], [
        ['知乎直答统一大脑', 'fast / thinking / agent 三档，服务端持有凭证，前端零配置'],
        ['四类 AI 角色', '辩手 / 对手 / 裁判 / 教练'],
        ['静默风格学习', '学习用户表达习惯，让 AI 越辩越"像你"，但立场始终对立'],
        ['七维裁判评分', '逻辑 / 论据 / 修辞 / 策略 / 表达 / 风度 / 伦理'],
        ['思考链与流式', '思考过程可折叠展示，发言打字机式流式输出'],
    ]),
    ('h2', '2.6 知识系统'),
    ('table', ['功能', '说明'], [
        ['本地 RAG', '纯前端中文分词 + BM25 + IndexedDB，支持 txt/md/docx/pdf 导入'],
        ['云端知识库', '服务端 FTS5 + BM25，共享资料与导入内容'],
        ['每日新闻学习', '8+ RSS 源自动抓取入库，辩论自动引用时事'],
        ['知乎源导入', '一键把回答 / 文章 / 收藏 / 收藏夹导入知识库'],
        ['11 个内置领域库', '开箱即用的领域知识'],
    ]),
    ('h2', '2.7 社区与成长系统'),
    ('table', ['功能', '说明'], [
        ['观点广场', '发帖 / 评论 / 点赞，支持删除自己发布的内容'],
        ['辩论报告', '每场自动生成结构化报告与专业建议'],
        ['个性化训练计划', '基于最弱维度生成 7 天可执行训练计划'],
        ['实力分匹配', 'Glicko-2 风格实力分 + 队列匹配 + 新手保护'],
        ['排行榜', '实力分 / 宠物双榜'],
        ['在线状态', '实时在线人数与状态展示'],
    ]),
    ('h2', '2.8 多端与交互系统'),
    ('table', ['功能', '说明'], [
        ['多端交付', 'Web（Vite）/ Android（Capacitor）/ Windows（Electron）'],
        ['语音输入', 'Web Speech API 实时转写，支持连续识别与中间结果'],
        ['麦克风解锁', 'HTTPS 3443 + CA 证书下载，手机信任后即可用'],
        ['全局按键音效', 'Web Audio 合成点击音，设置页可开关'],
        ['视觉体系', '深空星系设计语言 + FUI 人格观测站'],
    ]),
    ('pagebreak',),

    # ---------------- 三、创新能力 ----------------
    ('h1', '三、创新能力'),
    ('p', '本章是产品的差异化核心。创新点按"模式—体验—技术"三层归类，每一点都对应已落地的具体实现，而非概念规划。'),
    ('h2', '3.1 创新总览'),
    ('table', ['层级', '创新点', '差异化程度'], [
        ['模式创新', '人格即身份：MBTI 从"测试结果"变为"参与动机系统"', '高'],
        ['模式创新', '辩论即内容：实时交锋转为结构化可传播内容', '高'],
        ['模式创新', 'AI 对手订阅：把"找不到人辩"变成"永远有人陪练"', '高'],
        ['体验创新', '发言节奏引擎：逐句连发、真人感对话', '极高'],
        ['体验创新', '四场景玩法引擎：同一内核，四种气质', '高'],
        ['体验创新', '静默风格学习：AI 越辩越像你，立场却始终对立', '极高'],
        ['体验创新', '发言即攻击：抽象评分变成即时可视化战斗', '中高'],
        ['体验创新', '三源知识引擎：用户收藏什么，AI 就学什么', '中高'],
        ['技术创新', '知乎直答免配置统一大脑 + 多级降级', '中高'],
        ['技术创新', '服务器权威实时对战 + AI 影子用户同轨', '中'],
        ['技术创新', '全自动云端部署链路（分段上传—拼接—重启—自检）', '中'],
    ]),
    ('h2', '3.2 模式创新'),
    ('h3', '创新一 · 人格即身份'),
    ('p', '行业里 MBTI 普遍被当作"一次性测试结果"或"社交标签"，用户测完即走。本产品把它升级为参与动机系统：人格决定用户的入场身份、配色、头像光晕、发言立场与成长路径，并随辩论历史持续演化。人格不再是终点，而是起点。'),
    ('h3', '创新二 · 辩论即内容'),
    ('p', '实时辩论类产品普遍"爽完即走"，不留内容。本产品把一场辩论的全部产物（立论、反驳、总结、评分、金句）结构化落库，自动生成辩论报告并可发布为观点内容——让"消耗型娱乐"变成"生产型内容"。'),
    ('h3', '创新三 · AI 对手订阅制'),
    ('p', '辩论、对谈类产品最大的死穴是"找不到人"。本产品用 AI 补齐对手供给侧：AI 永远在线、三档难度、且会学习用户的表达习惯。用户不必等对手，也不必迁就对手水平——这使产品具备了"随时开局"的能力。'),
    ('h2', '3.3 体验创新'),
    ('h3', '创新四 · 发言节奏引擎（最强体验差异）'),
    ('p', '这是"像不像真人"的分水岭。多数 AI 对话产品把整段回复一次性倾倒，暴露"机器感"。本产品把 AI 回复拆成句子，以约 1.1 秒的节奏逐句发出，观点表达完整后再交还回合、切换阶段；配合 AI 主动开场与"请 AI 继续说"，形成接近真人对辩的呼吸感。'),
    ('h3', '创新五 · 四场景玩法引擎'),
    ('p', '同一套对战内核，通过"模式（对抗/自由）+ 目标 + 话题池 + 玩法提示"四要素，派生出法律、圆桌、脱口秀、街头四种截然不同的气质。架构上可复用，体验上不重复——这是可持续扩展玩法的引擎化设计。'),
    ('h3', '创新六 · 静默风格学习'),
    ('p', '系统在后台提取用户的表达特征（句长、反问频率、分点习惯、推进词、Emoji），生成"表达画像"注入 AI 提示词。效果是：AI 越辩越"像你"，句子长度与节奏与你同频——但立场始终是反方，保证对抗性不衰减。这是一项"用户感知不到、但一定能感觉到"的创新。'),
    ('h3', '创新七 · 发言即攻击的宠物可视化'),
    ('p', '把抽象的辩论评分，转化为直观的宠物战斗：发言即攻击，伤害由发言质量、逻辑、论据等共同决定，并由服务端权威计算、双端一致渲染。这为"辩论"这件认知性活动，叠加了一层即时反馈的游戏化外壳。'),
    ('h3', '创新八 · 三源知识引擎'),
    ('p', '本地 RAG（个人文档）+ 云端知识库（共享）+ 知乎源（用户内容与收藏）三源合一。用户把自己的知乎收藏导入后，AI 下一场就能引用其论点参与辩论——"把你自己的思想变成辩论武器"，形成了从内容消费到内容再生产的闭环。'),
    ('h2', '3.4 技术创新'),
    ('h3', '创新九 · 知乎直答免配置统一大脑'),
    ('p', '全站 AI 统一走知乎直答，凭证仅存服务端、前端零配置即可使用，用户不必自备 LLM Key。配合多级容错：thinking 档超时自动降级 fast、429 限流退避重试、无凭证时优雅降级——保证"永远不会因为 AI 挂了而整场崩"。'),
    ('h3', '创新十 · 服务器权威实时对战 + AI 影子用户'),
    ('p', '所有关键计算（伤害、评分、胜负）由服务端完成并广播，客户端只做渲染，从架构上杜绝作弊与不同步。AI 以影子用户身份与真人处于同一条发言时间线上，人机混战对前端完全透明。'),
    ('h3', '创新十一 · 全自动云端部署链路'),
    ('p', '后端内置自动化发布能力：整包 → 分段上传 → 服务器拼接 → 解压 → 重启 → 自检。绕开了受限端口与网络策略，实现"一条命令完成整包发布并验证"，显著降低了小团队的运维成本。'),
    ('h2', '3.5 与同类产品的创新对比'),
    ('table', ['能力', '通用 AI 聊天', '论坛 / 社区', '语音房', '思辩星球'], [
        ['人格化身份', '弱', '无', '无', '强（16 型 + 演化）'],
        ['AI 真实感发言', '中（整段）', '—', '—', '强（逐句连发）'],
        ['结构化对战', '无', '无', '弱', '强（三阶段 + 评分）'],
        ['内容沉淀', '无', '强', '无', '强（报告 + 广场）'],
        ['对手供给', '强', '弱', '不稳定', '强（AI 永远在线）'],
        ['个性化成长', '弱', '无', '无', '强（训练计划）'],
    ]),
    ('pagebreak',),

    # ---------------- 四、技术方案 ----------------
    ('h1', '四、技术方案'),
    ('h2', '4.1 总体架构'),
    ('table', ['层级', '组成', '职责'], [
        ['客户端层', 'Web（Vite/React18）、Android（Capacitor）、Windows（Electron）', '三端同源，共享业务逻辑与视觉系统'],
        ['接入层', 'Express + Socket.IO；HTTP 3001 / HTTPS 3443（自签 CA 可下载）', '路由、实时广播、麦克风解锁'],
        ['应用层', '人格引擎 / 辩论编排 / PK 房间 / 裁判评分 / 知识检索 / 社区', '承载全部业务逻辑，服务器权威'],
        ['AI 能力层', '知乎直答（fast / thinking / agent）+ 全网搜索 + 热榜 + 用户内容', '统一 AI 大脑与真实资料源'],
        ['数据层', 'SQLite（better-sqlite3, WAL）+ FTS5 + 前端 IndexedDB', '持久化与检索'],
        ['基础设施', '阿里云 ECS（cn-hangzhou）+ systemd + 自动化部署链路', '稳定运行与持续交付'],
    ]),
    ('h2', '4.2 关键技术实现'),
    ('h3', '① 发言节奏引擎'),
    ('p', 'AI 生成的整段文本 → 按换行与句末标点切分为句 → 循环以约 1.1 秒间隔写入发言表并广播 → 全部表达完毕后归还回合、推进阶段。配合"本阶段是否已发言"的判定，实现开场（kickoff）、追问（poke）、强制（force）三种触发；倒计时到点而 AI 仍在发言时自动切阶段并顺延。'),
    ('h3', '② 服务器权威与影子用户'),
    ('p', '评分与战斗结算在服务端完成；AI 以 ai__ 前缀的影子用户身份写入同一条发言流，与真人用户共享回合与阶段语义。发言队列带 busy 标志防止重入，避免 AI 重复开场或并发发言。'),
    ('h3', '③ 静默风格学习'),
    ('p', '每轮发言提取统计特征（平均句长、问句率、感叹率、分点率、推进词率、Emoji 率），累计到用户风格表；样本 ≥ 2 条后生成"表达画像"文本，注入中级 / 大师档 AI 的提示词。样本不足时不启用，避免过度拟合。'),
    ('h3', '④ 容错与降级'),
    ('p', 'AI 调用带 15 秒超时；thinking 档失败自动降级 fast 快速回；429 限流短暂退避后重试一次；未配置凭证时相关接口返回 503 并给出明确提示，前端优雅降级显示而非静默失败。'),
    ('h2', '4.3 技术栈'),
    ('table', ['领域', '选型'], [
        ['前端', 'React 18 · TypeScript · Vite · Tailwind · framer-motion'],
        ['服务端', 'Node.js · Express · better-sqlite3（WAL）· Socket.IO'],
        ['AI 能力', '知乎直答大模型（三档）· 全网搜索 · 热榜'],
        ['检索', '服务端 FTS5 + BM25；前端分词 + BM25 + IndexedDB'],
        ['多端', 'Capacitor（Android）· Electron（Windows）'],
        ['部署', '阿里云 ECS · systemd · Cloud Assistant 自动化'],
    ]),
    ('h2', '4.4 数据模型（核心表）'),
    ('table', ['数据域', '核心表'], [
        ['账号', 'users'],
        ['辩论会话', 'sessions / messages'],
        ['社区', 'posts / comments / likes'],
        ['人格库', 'personality_profiles / few_shot_examples'],
        ['实时对战', 'pk_rooms / pk_moves'],
        ['风格学习', 'ai_learner_styles'],
        ['成长与实力', 'rating / pet_dim_history / pet_skills'],
    ]),
    ('h2', '4.5 安全与合规'),
    ('bullets', [
        '凭证隔离 —— 知乎 Access Secret 仅存服务端 .env，前端永不接触。',
        '服务端代理 —— 第三方调用统一经服务端转发，Bearer 鉴权 + 时间戳 + RPM 限流。',
        '权限校验 —— JWT 鉴权；删除接口校验 user_id，非本人 403；用户只能操作自己的内容。',
        '降级透明 —— 未配置凭证时返回 503 + 明确提示，不静默失败。',
        '敏感文件不入库 —— 数据库、证书、签名密钥库均在 .gitignore 中。',
    ]),
    ('h2', '4.6 部署与运维'),
    ('p', '生产环境部署于阿里云 ECS（cn-hangzhou），systemd 常驻，同时开放 HTTP 3001 与 HTTPS 3443。后端内置自动化部署链路，支持整包发布并自检。'),
    ('h2', '4.7 关键技术取舍'),
    ('table', ['决策', '取舍', '理由'], [
        ['SQLite 而非 PostgreSQL', '牺牲并发上限', '单机轻量、零运维，WAL 足以支撑验证期规模'],
        ['自签 HTTPS 而非公网证书', '需引导信任证书', '内测期以最低成本快速解锁麦克风'],
        ['逐句连发而非整段', '单场时长略增', '真实感优先，是核心体验差异点'],
        ['缺凭证即降级 503', '功能不可用', '透明可诊断，降低排查成本'],
        ['服务器权威计算', '服务端压力略增', '保证多端一致与防作弊'],
    ]),
    ('pagebreak',),

    # ---------------- 五、知乎生态契合度 ----------------
    ('h1', '五、与知乎生态的契合度'),
    ('h2', '5.1 契合逻辑：同源的价值主张'),
    ('table', ['对标', '价值主张', '本质'], [
        ['知乎', '"有问题，就会有答案"', '知识的沉淀与分享'],
        ['思辩星球', '"有分歧，就会有交锋"', '观点的碰撞与激发'],
    ]),
    ('p', '二者不是竞争关系，而是上下游：知乎沉淀结论，本产品生产交锋。辩论产出的高质量观点——有立场、有论据、有反驳——恰恰是知乎最稀缺的"有理有据的回答"。'),
    ('h2', '5.2 能力契合：已落地的知乎能力'),
    ('table', ['知乎能力', '平台落地', '价值'], [
        ['直答大模型', '全站 AI 大脑（辩手 / 对手 / 裁判 / 教练）', '免配置接入中文一流 LLM'],
        ['全网搜索', '辩论大师"赛前资料包"', '论据从骨架升级为真实数据'],
        ['热榜', '每日热问 + 一键转辩题', '内容与热点实时同步'],
        ['用户内容', '一键导入知识库', '把"我的回答"变成 AI 素材'],
        ['用户收藏', '批量导入知识库', '收藏什么，AI 就学什么'],
        ['收藏夹', '按类目批量导入', '结构化知识管理'],
    ]),
    ('callout', '接入原则：凭证永不出服务端，前端只与自家代理对话；未配置时全链路优雅降级、不破坏任何既有功能。'),
    ('h2', '5.3 内容形态的天然映射'),
    ('table', ['知乎内容形态', '思辩星球对应物', '关系'], [
        ['回答', '一篇辩论报告 / 一次立论', '同构'],
        ['想法', '街头对谈 / 脱口秀金句', '同构'],
        ['圆桌', '圆桌会谈场景', '同名同构'],
        ['热榜', 'AI 实时命题来源', '上游'],
        ['收藏夹', 'AI 知识库语料', '上游'],
    ]),
    ('h2', '5.4 用户画像重合'),
    ('p', '知乎核心用户画像——高学历、爱表达、重逻辑、认同"理性讨论"——与本产品"表达型 / 逻辑型 / 沉思型"三类人群高度重合。用户从知乎带着"我想说点什么"的动机来到本平台，交锋后再把结论带回知乎。'),
    ('h2', '5.5 双向闭环'),
    ('p', '知乎热榜 / 问题 → AI 生成辩题 → 用户在平台实时交锋 → 高质量内容沉淀为辩论报告 → 以知乎式内容形态回流（回答 / 想法）→ 反哺知乎话题热度。整条链路让两个平台互相供给"话题"与"内容"。'),
    ('h2', '5.6 契合度自评'),
    ('table', ['维度', '契合度', '说明'], [
        ['价值观', '★★★★★  5/5', '"理性讨论、观点交锋"高度同源'],
        ['AI 能力', '★★★★★  5/5', '直答三档已实测接入并驱动全站 AI'],
        ['技术接口', '★★★★★  5/5', '已打通 6 类知乎 API，纯加法、无侵入'],
        ['内容形态', '★★★★☆  4/5', '回答/想法/圆桌天然映射，视频形态待补'],
        ['用户画像', '★★★★☆  4/5', '核心人群重合度高，需扩大覆盖'],
        ['商业化', '★★★☆☆  3/5', '合作路径清晰，变现模型待验证'],
        ['综合', '★★★★☆  4.3 / 5', '生态契合度整体高，技术侧已具备落地条件'],
    ]),
    ('h2', '5.7 合作路径'),
    ('bullets', [
        '账号互通 —— 知乎账号一键登录，人格档案与身份同步。',
        '内容双发 —— 辩论报告一键发布为知乎回答 / 想法，带平台署名。',
        '官方大脑 —— 知乎直答作为"官方推荐辩手"入驻。',
        '联合运营 —— 基于知乎热榜的辩题挑战赛、圆桌共创。',
        '数据反哺 —— 优质辩论观点结构化回流知乎问题页，抬高问答深度。',
    ]),
    ('pagebreak',),

    # ---------------- 结语 ----------------
    ('h1', '结语'),
    ('p', '思辩星球要做的，不是又一个聊天机器人，也不是又一个论坛。它试图回答一个更朴素的问题：当一个人有话想说时，能不能有一个地方，让他不必鼓起勇气、不必等待对手、不必担心白说——说完之后，还能看见自己变强了一点。'),
    ('callout', '人格给他身份，AI 给他对手，辩论给他内容，报告给他成长，而知乎生态给他更大的回响。'),
]


# ============================================================
# Word renderer
# ============================================================
def _rfonts(run, ascii_font, ea_font):
    rPr = run._element.get_or_add_rPr()
    rf = rPr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rPr.insert(0, rf)
    rf.set(qn('w:ascii'), ascii_font); rf.set(qn('w:hAnsi'), ascii_font)
    rf.set(qn('w:eastAsia'), ea_font); rf.set(qn('w:cs'), ascii_font)


def style_run(run, size=10.5, bold=False, color=INK, italic=False, font=EN, ea=CN):
    run.font.size = Pt(size); run.bold = bold; run.italic = italic
    run.font.color.rgb = color
    _rfonts(run, font, ea)


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def set_cell_margins(table, top=60, bottom=60, left=100, right=100):
    mar = OxmlElement('w:tblCellMar')
    for tag, val in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        el = OxmlElement('w:' + tag); el.set(qn('w:w'), str(val)); el.set(qn('w:type'), 'dxa'); mar.append(el)
    table._tbl.tblPr.append(mar)


def table_borders(table, color=BORDER, sz=6):
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), str(sz))
        el.set(qn('w:space'), '0'); el.set(qn('w:color'), color)
        borders.append(el)
    table._tbl.tblPr.append(borders)


def para_spacing(p, before=0, after=6, line=1.5):
    pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after); pf.line_spacing = line
    return p


def add_page_number(paragraph):
    run = paragraph.add_run()
    style_run(run, size=8.5, color=MUTED)
    f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin')
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = 'PAGE'
    f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'end')
    run._r.append(f1); run._r.append(it); run._r.append(f2)


def build_docx():
    doc = Document()
    normal = doc.styles['Normal']
    normal.font.name = EN; normal.font.size = Pt(10.5); normal.font.color.rgb = INK
    rpr = normal.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
    rf.set(qn('w:ascii'), EN); rf.set(qn('w:hAnsi'), EN); rf.set(qn('w:eastAsia'), CN)

    sec = doc.sections[0]
    sec.top_margin = Cm(2.2); sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.2); sec.right_margin = Cm(2.2)

    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(fp)

    first = {'v': True}
    for block in CONTENT:
        kind = block[0]

        if kind == 'cover':
            p = doc.add_paragraph(); para_spacing(p, before=90, after=4, line=1.0)
            style_run(p.add_run('◉ 思辩星球  DebateSphere'), size=12, bold=True, color=PRIMARY)
            p = doc.add_paragraph(); para_spacing(p, before=40, after=2, line=1.0)
            style_run(p.add_run('产品方案与创新说明书'), size=32, bold=True, color=INK)
            p = doc.add_paragraph(); para_spacing(p, before=0, after=10, line=1.2)
            style_run(p.add_run('核心思路 · 功能架构 · 创新能力 · 技术方案 · 知乎生态契合度'), size=13, color=ACCENT2)
            p = doc.add_paragraph(); para_spacing(p, before=2, after=14, line=1.0)
            style_run(p.add_run('─' * 34), size=10, color=PRIMARY)
            p = doc.add_paragraph(); para_spacing(p, before=0, after=4, line=1.3)
            style_run(p.add_run('让每一种人格，都成为一颗星球。'), size=12, italic=True, color=INK2)
            p = doc.add_paragraph(); para_spacing(p, before=120, after=0, line=1.6)
            for line in ['文档版本：v1.0', '编制日期：2026 年 9 月', '文档密级：内部资料']:
                style_run(p.add_run(line + '\n'), size=10, color=MUTED)
            doc.add_page_break()
            continue

        if kind == 'pagebreak':
            doc.add_page_break(); first['v'] = True; continue

        if kind == 'h1':
            if not first['v']:
                doc.add_page_break()
            first['v'] = False
            p = doc.add_paragraph(); para_spacing(p, before=0, after=10, line=1.15)
            style_run(p.add_run(block[1]), size=19, bold=True, color=INK)
            pPr = p._p.get_or_add_pPr(); pbdr = OxmlElement('w:pBdr')
            b = OxmlElement('w:bottom')
            b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '12'); b.set(qn('w:space'), '6'); b.set(qn('w:color'), H1_LINE)
            pbdr.append(b); pPr.append(pbdr)
            continue

        if kind == 'h2':
            p = doc.add_paragraph(); para_spacing(p, before=12, after=5, line=1.2)
            style_run(p.add_run(block[1]), size=14, bold=True, color=PRIMARY)
            continue

        if kind == 'h3':
            p = doc.add_paragraph(); para_spacing(p, before=8, after=4, line=1.2)
            style_run(p.add_run(block[1]), size=12, bold=True, color=INK)
            continue

        if kind == 'p':
            p = doc.add_paragraph(); para_spacing(p, before=0, after=7, line=1.55)
            style_run(p.add_run(block[1]), size=10.5, color=INK2)
            continue

        if kind in ('bullets', 'nums'):
            for i, item in enumerate(block[1]):
                p = doc.add_paragraph(); para_spacing(p, before=0, after=4, line=1.5)
                p.paragraph_format.left_indent = Cm(0.75)
                p.paragraph_format.first_line_indent = Cm(-0.42)
                style_run(p.add_run('•  ' if kind == 'bullets' else '%d.  ' % (i + 1)), size=10.5, bold=True, color=PRIMARY)
                style_run(p.add_run(item), size=10.5, color=INK2)
            continue

        if kind == 'quote':
            p = doc.add_paragraph(); para_spacing(p, before=4, after=8, line=1.5)
            p.paragraph_format.left_indent = Cm(0.5)
            pPr = p._p.get_or_add_pPr(); pbdr = OxmlElement('w:pBdr')
            l = OxmlElement('w:left')
            l.set(qn('w:val'), 'single'); l.set(qn('w:sz'), '18'); l.set(qn('w:space'), '8'); l.set(qn('w:color'), H1_LINE)
            pbdr.append(l); pPr.append(pbdr)
            style_run(p.add_run(block[1]), size=10.5, italic=True, color=INK2)
            continue

        if kind == 'callout':
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            table_borders(tbl, color=CALLOUT_BORDER, sz=4)
            set_cell_margins(tbl, 100, 100, 160, 140)
            c = tbl.cell(0, 0); shade_cell(c, CALLOUT_FILL); c.text = ''
            p = c.paragraphs[0]; para_spacing(p, before=0, after=0, line=1.45)
            style_run(p.add_run(block[1]), size=10.5, bold=True, color=CALLOUT_TEXT)
            sp = doc.add_paragraph(); para_spacing(sp, before=0, after=0, line=1.0)
            continue

        if kind == 'table':
            headers, rows = block[1], block[2]
            ncol = len(headers)
            tbl = doc.add_table(rows=1, cols=ncol)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            table_borders(tbl); set_cell_margins(tbl)
            for j, htext in enumerate(headers):
                c = tbl.cell(0, j); shade_cell(c, H_FILL); c.text = ''
                p = c.paragraphs[0]; para_spacing(p, before=1, after=1, line=1.2)
                style_run(p.add_run(htext), size=9.5, bold=True, color=WHITE)
            for i, row in enumerate(rows):
                cells = tbl.add_row().cells
                for j in range(ncol):
                    c = cells[j]
                    if i % 2 == 1:
                        shade_cell(c, Z_FILL)
                    c.text = ''
                    p = c.paragraphs[0]; para_spacing(p, before=1, after=1, line=1.3)
                    style_run(p.add_run(row[j] if j < len(row) else ''), size=9.5, color=INK2, bold=(j == 0))
            sp = doc.add_paragraph(); para_spacing(sp, before=0, after=2, line=1.0)
            continue

    doc.save(DOCX_PATH)
    return DOCX_PATH


# ============================================================
# HTML renderer
# ============================================================
def esc(t):
    return H.escape(t)


def build_html():
    css = """
    @page { size: A4; margin: 18mm 16mm; }
    * { box-sizing: border-box; }
    body { margin:0; background:#eef1f6; color:#1e293b;
      font-family:"Segoe UI","Microsoft YaHei","微软雅黑",-apple-system,sans-serif;
      font-size:14px; line-height:1.7; }
    .page { max-width:820px; margin:0 auto; background:#fff; padding:56px 60px 72px;
      box-shadow:0 10px 40px rgba(15,23,42,.10); }
    .cover { min-height:960px; display:flex; flex-direction:column; justify-content:center; padding:60px 0; }
    .brand { color:#0f766e; font-weight:700; font-size:15px; letter-spacing:.5px; }
    h1.title { font-size:42px; font-weight:800; letter-spacing:-1px; margin:26px 0 10px; color:#0f172a; }
    .subtitle { font-size:16px; color:#b45309; font-weight:600; margin-bottom:18px; }
    .rule { height:4px; width:120px; background:linear-gradient(90deg,#0f766e,#b45309); border-radius:3px; margin:0 0 22px; }
    .tagline { font-size:17px; color:#475569; font-style:italic; }
    .meta { margin-top:120px; color:#64748b; font-size:13px; line-height:2; }
    h1 { font-size:25px; font-weight:800; color:#0f172a; margin:44px 0 16px; padding-bottom:10px;
      border-bottom:2.5px solid #0f766e; letter-spacing:-.4px; page-break-after:avoid; }
    h2 { font-size:18px; font-weight:700; color:#0f766e; margin:28px 0 10px; page-break-after:avoid; }
    h3 { font-size:15px; font-weight:700; color:#1e293b; margin:20px 0 8px; page-break-after:avoid; }
    p { margin:0 0 12px; color:#475569; }
    ul,ol { margin:0 0 14px; padding-left:22px; }
    li { margin-bottom:6px; color:#475569; }
    blockquote { margin:14px 0; padding:6px 0 6px 16px; border-left:3px solid #0f766e;
      color:#475569; font-style:italic; }
    .callout { background:#fffbeb; border:1px solid #fde68a; border-radius:10px; padding:14px 18px;
      margin:14px 0; color:#92400e; font-weight:600; }
    table { width:100%; border-collapse:collapse; margin:14px 0 20px; font-size:13px; page-break-inside:avoid; }
    th { background:#134e4a; color:#fff; text-align:left; padding:9px 12px; font-weight:600; font-size:12.5px; }
    td { padding:8px 12px; border-bottom:1px solid #cbd5e1; color:#475569; vertical-align:top; }
    td:first-child { color:#1e293b; font-weight:600; }
    tbody tr:nth-child(even) td { background:#f0fdfa; }
    .sec { page-break-before:always; }
    .foot { margin-top:40px; padding-top:14px; border-top:1px solid #e2e8f0;
      color:#94a3b8; font-size:12px; text-align:center; }
    @media print { body{background:#fff;} .page{box-shadow:none; margin:0; max-width:none; padding:0;} }
    """
    out = ['<!DOCTYPE html><html lang="zh-CN"><head><meta charset="utf-8">',
           '<meta name="viewport" content="width=device-width, initial-scale=1">',
           '<title>产品方案与创新说明书 · 思辩星球 DebateSphere</title>',
           '<style>%s</style></head><body><div class="page">' % css]

    for block in CONTENT:
        k = block[0]
        if k == 'cover':
            out.append('<div class="cover">'
                       '<div class="brand">◉ 思辩星球 &nbsp;DebateSphere</div>'
                       '<h1 class="title">产品方案与创新说明书</h1>'
                       '<div class="subtitle">核心思路 · 功能架构 · 创新能力 · 技术方案 · 知乎生态契合度</div>'
                       '<div class="rule"></div>'
                       '<div class="tagline">让每一种人格，都成为一颗星球。</div>'
                       '<div class="meta">文档版本：v1.0<br>编制日期：2026 年 9 月<br>文档密级：内部资料</div>'
                       '</div>')
        elif k == 'pagebreak':
            out.append('<div class="sec"></div>')
        elif k == 'h1':
            out.append('<h1>%s</h1>' % esc(block[1]))
        elif k == 'h2':
            out.append('<h2>%s</h2>' % esc(block[1]))
        elif k == 'h3':
            out.append('<h3>%s</h3>' % esc(block[1]))
        elif k == 'p':
            out.append('<p>%s</p>' % esc(block[1]))
        elif k == 'bullets':
            out.append('<ul>' + ''.join('<li>%s</li>' % esc(x) for x in block[1]) + '</ul>')
        elif k == 'nums':
            out.append('<ol>' + ''.join('<li>%s</li>' % esc(x) for x in block[1]) + '</ol>')
        elif k == 'quote':
            out.append('<blockquote>%s</blockquote>' % esc(block[1]))
        elif k == 'callout':
            out.append('<div class="callout">%s</div>' % esc(block[1]))
        elif k == 'table':
            headers, rows = block[1], block[2]
            th = ''.join('<th>%s</th>' % esc(x) for x in headers)
            tb = ''.join('<tr>' + ''.join('<td>%s</td>' % esc(c) for c in r) + '</tr>' for r in rows)
            out.append('<table><thead><tr>%s</tr></thead><tbody>%s</tbody></table>' % (th, tb))

    out.append('<div class="foot">思辩星球 DebateSphere · 产品方案与创新说明书 v1.0 · 2026-09 · 内部资料</div>')
    out.append('</div></body></html>')

    with io.open(HTML_PATH, 'w', encoding='utf-8') as f:
        f.write('\n'.join(out))
    return HTML_PATH


if __name__ == '__main__':
    d = build_docx(); h = build_html()
    print('DOCX:', d, os.path.getsize(d), 'bytes')
    print('HTML:', h, os.path.getsize(h), 'bytes')
